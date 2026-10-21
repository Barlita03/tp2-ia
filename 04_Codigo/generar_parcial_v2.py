"""
Genera la Entrega Parcial - v2 con todos los datos reales y correcciones
"""

import os
import shutil
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    root_dir = os.path.dirname(script_dir)
    docs_orig_dir = os.path.join(root_dir, "02_Documentos_Originales")
    entregables_dir = os.path.join(root_dir, "01_Entregables")
    ejemplos_dir = os.path.join(root_dir, "05_Resultados_Graficos", "ejemplos_informe")

    src_path = os.path.join(docs_orig_dir, "Entrega Parcial - TP°2 - Sistemas Inteligentes - Grupo 9 - Hongos.docx")
    dst_path = os.path.join(entregables_dir, "Entrega Parcial - TP°2 - Sistemas Inteligentes - Grupo 9 - Hongos - v2.docx")

    shutil.copyfile(src_path, dst_path)
    doc = docx.Document(dst_path)

    for p in doc.paragraphs:
        text = p.text

        # 1. Correccion indispensable aprendizaje profundo y riesgo anexo
        if "resulta indispensable el uso de modelos de aprendizaje profundo" in text:
            p.text = (
                "Estas características espaciales y de textura complejas impiden resolver el problema con programación "
                "tradicional o reglas deterministas, por lo que resulta adecuado el uso de Redes Neuronales Artificiales "
                "basadas en visión por computadora, en particular arquitecturas convolucionales (CNN) combinadas con un modelo "
                "base de Perceptrón Multicapa (MLP).\n\n"
                "Aclaración metodológica sobre el Anexo de problemas resueltos: Se destaca que el problema aquí abordado difiere "
                "sustancialmente del caso consignado en el Anexo de la cátedra ('Predecir si el hongo es venenoso o no a partir "
                "de las características visibles'), el cual refiere a una clasificación binaria clásica (comestible vs. venenoso) sobre "
                "datos tabulares pre-categorizados en laboratorio (dataset histórico de UCI de 22 atributos nominales). En cambio, el "
                "presente trabajo aborda una tarea de visión por computadora multiclase sobre fotografías no estructuradas en entorno "
                "natural, orientada a la identificación taxonómica fina de especies."
            )

        # 2. Correccion filtros en CNN Modelo 3
        elif "reduciendo suavemente la cantidad de filtros a medida que aumenta la profundidad" in text:
            p.text = (
                "Modelo 3 (Arquitectura profunda de dos capas convolucionales): una ConvNet con dos capas convolucionales sucesivas "
                "(Conv1 y Conv2), cada una con su etapa de activación no lineal (ReLU) y submuestreo espacial (Max Pooling), culminando "
                "en una capa de aplanado (Flatten/Reshape) y capas lineales densas con función de salida Softmax. De acuerdo con las "
                "pautas pedagógicas de la materia y la guía de estudio (Módulo 37), la cantidad de filtros debe AUMENTAR con la "
                "profundidad (por ejemplo, 32 filtros en Conv1 y 64 filtros en Conv2 en potencias de 2). De esta forma se compensa la "
                "reducción progresiva de la resolución espacial con un incremento en la riqueza de representaciones y combinaciones "
                "de texturas complejas, evitando cuellos de botella en la extracción jerárquica de características."
            )

        # 3. Cantidad de imagenes y clases reales
        elif "[completar tras descargar el dataset]" in text:
            p.text = (
                "Fuente de datos: dataset público de Kaggle 'Mushroom Classification' (autor: tinkhoav), orientado a visión por "
                "computadora. Cantidad total de imágenes: 83.729 imágenes distribuidas en 277 clases (especies de hongos). Tras el análisis "
                "exploratorio de integridad y el cálculo de hashes criptográficos MD5, se detectaron 15 instancias duplicadas exactas "
                "(14 intra-clase y 1 caso de conflicto inter-clase entre Collybia tuberosa y Lactarius helvus, descartado por inconsistencia). "
                "El conjunto limpio consta de 83.714 imágenes únicas. La distribución presenta un desbalance medido con un mínimo de 24 imágenes "
                "por clase, un máximo de 444, media de 302,27 y mediana de 325 imágenes (ratio de desbalance de 18,50x), lo que fundamenta el uso "
                "de pesos balanceados por clase (class weights) y el reporte de la métrica F1-score macro junto con la exactitud."
            )

        # 4. Resolucion y normalizacion
        elif "archivos JPG o PNG), que serán reescaladas a una dimensión estándar" in text:
            p.text = (
                "Atributos de entrada: imágenes digitales a color en formato RGB (archivos JPG/JPEG de 3 canales). Todas las imágenes "
                "del repositorio cuentan con resolución original uniforme de 224 x 224 píxeles. Para optimizar el entrenamiento y mitigar la "
                "explosión combinatoria de parámetros según el tipo de red, se adopta:\n"
                "• Para el Perceptrón Multicapa (MLP): redimensionamiento a 64 x 64 píxeles (12.288 entradas al vectorizar 64x64x3). Esto es "
                "crítico, ya que mantener 128x128 exigiría 49.152 entradas, desatando más de 6,2 millones de pesos sinápticos solo en la primera "
                "capa oculta, lo que confirma la advertencia teórica de la cátedra sobre la ineficiencia de los MLP para imágenes complejas.\n"
                "• Para las Redes Neuronales Convolucionales (Modelos 2 y 3): redimensionamiento a 128 x 128 píxeles (o 224 x 224), lo cual "
                "preserva con nitidez los detalles texturales del sombrero y las láminas del hongo con un número reducido de parámetros "
                "gracias a la compartición de pesos y campos receptivos locales.\n"
                "• Normalización: escalado lineal de los valores de píxel dividiendo por 255.0, proyectando las intensidades [0, 255] al rango "
                "continuo [0.0, 1.0], compatible con la función de activación ReLU y la estabilidad del gradiente."
            )

        # 5. Cantidades exactas de particion
        elif "Entrenamiento: 60% de las imágenes." in text:
            p.text = (
                "Entrenamiento: 60% de las imágenes (50.228 imágenes), sobre las cuales se ajustan los pesos de la red y se aplica data augmentation."
            )
        elif "Validación: 20% de las imágenes." in text:
            p.text = (
                "Validación: 20% de las imágenes (16.743 imágenes), reservadas para el seguimiento del error de generalización (loss), ajuste de tasa de aprendizaje y early stopping."
            )
        elif "Prueba (Test): 20% de las imágenes." in text:
            p.text = (
                "Prueba (Test): 20% de las imágenes (16.743 imágenes), destinadas exclusivamente a la evaluación final del modelo sobre datos nunca vistos durante el entrenamiento, dando estricto cumplimiento al requisito de la consigna de reservar como mínimo un 20% para prueba."
            )

        # 6. Placeholder de ejemplos de imagenes
        elif "[agregar 2 o 3 fotos representativas del dataset]" in text:
            p.text = (
                "Ejemplos de imágenes por clase del dataset: A continuación se presentan ejemplares de especies representativas extraídas "
                "directamente del dataset, evidenciando tanto la alta variabilidad morfológica como la sutil diferenciación visual requerida:"
            )

    # Agregar las imagenes al final del documento
    fotos = [
        ("Amanita_muscaria_toxica.jpg", "Figura 1: Amanita muscaria (Especie tóxica, llamativo sombrero rojo con motas blancas)."),
        ("Amanita_pantherina_toxica.jpg", "Figura 2: Amanita pantherina (Especie altamente tóxica, sombrero pardo oliváceo)."),
        ("Boletus_edulis_comestible.jpg", "Figura 3: Boletus edulis (Especie comestible silvestre muy apreciada)."),
        ("Agaricus_campestris_champiñon.jpg", "Figura 4: Agaricus campestris (Champiñón silvestre comestible de campo).")
    ]

    for fn, caption in fotos:
        fpath = os.path.join(ejemplos_dir, fn)
        if os.path.exists(fpath):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p_img.add_run()
            run.add_picture(fpath, width=Inches(2.8))
            p_cap = doc.add_paragraph(caption)
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.runs[0].font.size = Pt(9.5)
            p_cap.runs[0].font.italic = True

    doc.save(dst_path)
    print("Entrega Parcial v2 guardada exitosamente en:", dst_path)

if __name__ == "__main__":
    main()
